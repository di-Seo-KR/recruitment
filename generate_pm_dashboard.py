#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PM 업무관리 대시보드(엑셀) 생성 스크립트.

To Do / Doing / Done 칸반 방식의 일정관리 툴을 xlsx로 생성한다.
  - 사용법   : 사용 안내
  - 대시보드 : KPI 타일 + 상태/우선순위/담당자별 현황 + 차트 (전부 자동 계산)
  - 칸반보드 : To Do / Doing / Done 카드 보드 (자동 갱신)
  - 업무목록 : 실제 입력 시트 (드롭다운, 조건부서식, D-Day 자동 계산)
  - 설정     : 드롭다운 목록 관리

실행:  python3 generate_pm_dashboard.py  →  PM_업무관리_대시보드.xlsx
"""
import datetime as dt

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "PM_업무관리_대시보드.xlsx"
FONT = "맑은 고딕"
N_ROWS = 200                 # 업무 입력 가능 행 수 (2~201행)
LAST = N_ROWS + 1            # 마지막 데이터 행 번호
KANBAN_SLOTS = 30            # 칸반 열당 카드 수

# ── 색상 (dataviz 검증 팔레트) ────────────────────────────────────────────
INK = "0B0B0B"; INK2 = "52514E"; MUTED = "898781"
SURFACE = "FCFCFB"; PAGE = "F9F9F7"; HAIR = "E1E0D9"; BASE = "C3C2B7"
BLUE = "2A78D6"; BLUE_DK = "1C5CAB"; BLUE_LT = "CDE2FB"; BLUE_MD = "86B6EF"
GREEN = "0CA30C"; GREEN_TX = "006300"; GREEN_LT = "DFF2DF"; GREEN_MD = "A5D6A5"
RED = "D03B3B"; RED_LT = "F7DEDE"; RED_WASH = "FBEDEC"
AMBER = "C98500"; AMBER_LT = "FDF2D9"
GRAY_LT = "F0EFEC"

STATUS = ["To Do", "Doing", "Done"]
PRIORITY = ["높음", "중간", "낮음"]
MEMBERS = ["김민준", "이서연", "박지훈", "최수아", "정다은"]

thin = Side(style="thin", color=HAIR)
thin_bd = Border(left=thin, right=thin, top=thin, bottom=thin)


def f(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def box(ws, r1, c1, r2, c2, color=HAIR, style="thin"):
    """범위 바깥 테두리를 그린다(안쪽 하이라인 유지)."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            bd = cell.border
            s = {k: getattr(bd, k) for k in ("left", "right", "top", "bottom")}
            if r == r1:
                s["top"] = Side(style=style, color=color)
            if r == r2:
                s["bottom"] = Side(style=style, color=color)
            if c == c1:
                s["left"] = Side(style=style, color=color)
            if c == c2:
                s["right"] = Side(style=style, color=color)
            cell.border = Border(**s)


def paint(ws, r1, c1, r2, c2, hex_):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).fill = fill(hex_)


wb = Workbook()
ws_help = wb.active
ws_help.title = "사용법"
ws_dash = wb.create_sheet("대시보드")
ws_kanban = wb.create_sheet("칸반보드")
ws_task = wb.create_sheet("업무목록")
ws_cfg = wb.create_sheet("설정")

ws_help.sheet_properties.tabColor = MUTED
ws_dash.sheet_properties.tabColor = BLUE
ws_kanban.sheet_properties.tabColor = AMBER
ws_task.sheet_properties.tabColor = GREEN
ws_cfg.sheet_properties.tabColor = MUTED

# ═══════════════════════════════ 설정 ═══════════════════════════════
ws_cfg["A1"] = "⚙️ 설정 — 드롭다운 목록 관리"
ws_cfg["A1"].font = f(14, True)
ws_cfg["A2"] = "담당자만 수정하세요. 상태/우선순위는 수식과 연결되어 있어 바꾸면 안 됩니다."
ws_cfg["A2"].font = f(9, color=MUTED)

for col, (title, note) in zip(
    "ABC",
    [("상태", "변경 금지"), ("우선순위", "변경 금지"), ("담당자", "자유롭게 수정")],
):
    c = ws_cfg[f"{col}3"]
    c.value = f"{title} ({note})"
    c.font = f(10, True, "FFFFFF")
    c.fill = fill(INK2 if note == "변경 금지" else GREEN)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws_cfg.column_dimensions[col].width = 20

for i, v in enumerate(STATUS):
    ws_cfg.cell(row=4 + i, column=1, value=v)
for i, v in enumerate(PRIORITY):
    ws_cfg.cell(row=4 + i, column=2, value=v)
for i, v in enumerate(MEMBERS):
    ws_cfg.cell(row=4 + i, column=3, value=v)

for r in range(4, 14):
    for c in range(1, 4):
        if c < 3 and r > 6:
            continue
        cell = ws_cfg.cell(row=r, column=c)
        cell.font = f(10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_bd
        if c == 3:
            cell.fill = fill(GREEN_LT if not cell.value else "FFFFFF")

ws_cfg["A16"] = "※ 담당자는 최대 10명까지(C4~C13). 이름을 바꾸면 업무목록 드롭다운과 대시보드에 바로 반영됩니다."
ws_cfg["A16"].font = f(9, color=MUTED)

for name, ref in [
    ("StatusList", "설정!$A$4:$A$6"),
    ("PriorityList", "설정!$B$4:$B$6"),
    ("MemberList", "설정!$C$4:$C$13"),
]:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ═══════════════════════════════ 업무목록 ═══════════════════════════════
HEADERS = ["No.", "업무명", "프로젝트", "담당자", "우선순위", "상태",
           "시작일", "마감일", "진행률", "D-Day", "메모"]
WIDTHS = [6, 38, 15, 11, 10, 10, 12, 12, 10, 14, 28]

for i, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
    c = ws_task.cell(row=1, column=i, value=h)
    c.font = f(10, True, "FFFFFF")
    c.fill = fill(INK2)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = Border(bottom=Side(style="medium", color=INK))
    ws_task.column_dimensions[get_column_letter(i)].width = w
ws_task.row_dimensions[1].height = 26

# 숨김 도우미 열: L/M/N = 상태별 순번(칸반용), O = 카드 표시 텍스트
for col, h in zip("LMNO", ["_todo", "_doing", "_done", "_card"]):
    ws_task[f"{col}1"] = h
    ws_task[f"{col}1"].font = f(8, color=MUTED)
    ws_task.column_dimensions[col].hidden = True

today = dt.date.today()
SAMPLES = [
    ("신규 서비스 런칭 킥오프 미팅", "런칭 프로젝트", "김민준", "높음", "Done", -14, -10, 100, "회의록 공유 완료"),
    ("요구사항 정의서(PRD) 작성", "런칭 프로젝트", "김민준", "높음", "Done", -10, -3, 100, "v1.2 확정"),
    ("와이어프레임 시안 검토", "런칭 프로젝트", "이서연", "중간", "Doing", -5, 1, 70, "디자인팀 피드백 반영 중"),
    ("개발 일정 산정 및 스프린트 계획", "런칭 프로젝트", "박지훈", "높음", "Doing", -3, 3, 40, ""),
    ("경쟁사 기능 벤치마킹 리서치", "리서치", "최수아", "낮음", "Doing", -7, -1, 80, "지연 — 보고서 마무리 필요"),
    ("유저 인터뷰 대상자 모집", "리서치", "최수아", "중간", "To Do", 1, 7, 0, "10명 목표"),
    ("온보딩 플로우 개선안 정리", "UX 개선", "이서연", "중간", "To Do", 3, 10, 0, ""),
    ("3분기 로드맵 경영진 보고", "운영", "김민준", "높음", "To Do", 5, 9, 0, "발표자료 필요"),
    ("데이터 대시보드 지표 정의", "데이터", "박지훈", "중간", "Doing", -2, 5, 30, "핵심 지표 3개 선정"),
    ("고객 문의(VOC) 월간 리포트", "운영", "정다은", "낮음", "Done", -8, -2, 100, ""),
    ("A/B 테스트 결과 분석 및 공유", "데이터", "정다은", "중간", "To Do", 2, 8, 0, ""),
    ("개인정보 처리방침 개정 검토", "운영", "최수아", "높음", "To Do", 0, 6, 0, "법무 검토 병행"),
]

for i, (name, proj, owner, pri, st, s_off, d_off, prog, memo) in enumerate(SAMPLES):
    r = 2 + i
    ws_task.cell(row=r, column=2, value=name)
    ws_task.cell(row=r, column=3, value=proj)
    ws_task.cell(row=r, column=4, value=owner)
    ws_task.cell(row=r, column=5, value=pri)
    ws_task.cell(row=r, column=6, value=st)
    ws_task.cell(row=r, column=7, value=today + dt.timedelta(days=s_off))
    ws_task.cell(row=r, column=8, value=today + dt.timedelta(days=d_off))
    ws_task.cell(row=r, column=9, value=prog)
    ws_task.cell(row=r, column=11, value=memo)

for r in range(2, LAST + 1):
    ws_task.cell(row=r, column=1, value=f'=IF($B{r}="","",COUNTA($B$2:$B{r}))')
    ws_task.cell(row=r, column=10, value=(
        f'=IF($B{r}="","",IF($F{r}="Done","✔ 완료",IF($H{r}="","—",'
        f'IF($H{r}<TODAY(),"⚠ "&(TODAY()-$H{r})&"일 지연",'
        f'IF($H{r}=TODAY(),"🔔 오늘 마감","D-"&($H{r}-TODAY()))))))'))
    for col, st in zip("LMN", STATUS):
        ws_task[f"{col}{r}"] = (
            f'=IF(AND($B{r}<>"",$F{r}="{st}"),'
            f'COUNTIFS($B$2:$B{r},"<>",$F$2:$F{r},"{st}"),"")')
    ws_task[f"O{r}"] = (
        f'=IF($B{r}="","",$B{r}&"  |  "&IF($D{r}="","담당자 미지정",$D{r})'
        f'&IF($H{r}="",""," · ~"&TEXT($H{r},"m/d"))'
        f'&IF($E{r}="높음"," 🔺","")'
        f'&IF(AND($F{r}<>"Done",$H{r}<>"",$H{r}<TODAY())," ⚠",""))')

    for c in range(1, 12):
        cell = ws_task.cell(row=r, column=c)
        cell.font = f(10)
        cell.border = Border(bottom=thin)
        if c in (1, 4, 5, 6, 7, 8, 9, 10):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(vertical="center")
    ws_task.cell(row=r, column=7).number_format = "yyyy-mm-dd"
    ws_task.cell(row=r, column=8).number_format = "yyyy-mm-dd"
    ws_task.cell(row=r, column=9).number_format = '0"%"'

ws_task.freeze_panes = "C2"
ws_task.auto_filter.ref = f"A1:K{LAST}"

# 드롭다운
dv_status = DataValidation(type="list", formula1="StatusList", allow_blank=True,
                           errorTitle="상태 선택", error="To Do / Doing / Done 중에서 선택하세요.",
                           showErrorMessage=True)
dv_pri = DataValidation(type="list", formula1="PriorityList", allow_blank=True,
                        errorTitle="우선순위 선택", error="높음 / 중간 / 낮음 중에서 선택하세요.",
                        showErrorMessage=True)
dv_member = DataValidation(type="list", formula1="MemberList", allow_blank=True,
                           errorTitle="담당자 선택", error="설정 시트의 담당자 목록에서 선택하세요.",
                           showErrorMessage=True)
dv_date = DataValidation(type="date", operator="between", formula1="36526", formula2="73415",
                         allow_blank=True, errorTitle="날짜 입력",
                         error="날짜 형식으로 입력하세요. 예: 2026-07-10", showErrorMessage=True)
dv_prog = DataValidation(type="whole", operator="between", formula1="0", formula2="100",
                         allow_blank=True, errorTitle="진행률",
                         error="0~100 사이 숫자로 입력하세요.", showErrorMessage=True)
for dv, rng in [(dv_member, f"D2:D{LAST}"), (dv_pri, f"E2:E{LAST}"),
                (dv_status, f"F2:F{LAST}"), (dv_date, f"G2:H{LAST}"),
                (dv_prog, f"I2:I{LAST}")]:
    ws_task.add_data_validation(dv)
    dv.add(rng)

# 조건부 서식 (추가 순서 = 우선순위)
rng_F = f"F2:F{LAST}"
for val, fl, fc in [("To Do", GRAY_LT, INK2), ("Doing", BLUE_LT, BLUE_DK), ("Done", GREEN_LT, GREEN_TX)]:
    ws_task.conditional_formatting.add(
        rng_F, CellIsRule(operator="equal", formula=[f'"{val}"'],
                          fill=fill(fl), font=f(10, True, fc)))
rng_E = f"E2:E{LAST}"
for val, fl, fc, bold in [("높음", RED_LT, RED, True), ("중간", AMBER_LT, AMBER, False),
                          ("낮음", None, MUTED, False)]:
    kw = {"font": f(10, bold, fc)}
    if fl:
        kw["fill"] = fill(fl)
    ws_task.conditional_formatting.add(
        rng_E, CellIsRule(operator="equal", formula=[f'"{val}"'], **kw))

overdue = 'AND($B2<>"",$F2<>"Done",$H2<>"",$H2<TODAY())'
ws_task.conditional_formatting.add(
    f"J2:J{LAST}", FormulaRule(formula=[overdue], font=f(10, True, RED)))
ws_task.conditional_formatting.add(
    f"A2:K{LAST}", FormulaRule(formula=[overdue], fill=fill(RED_WASH)))
ws_task.conditional_formatting.add(
    f"A2:K{LAST}", FormulaRule(formula=['AND($B2<>"",$F2="Done")'], font=f(10, color=MUTED)))
ws_task.conditional_formatting.add(
    f"I2:I{LAST}", DataBarRule(start_type="num", start_value=0,
                               end_type="num", end_value=100,
                               color=BLUE, showValue=True))

# ═══════════════════════════════ 칸반보드 ═══════════════════════════════
ws_kanban.sheet_view.showGridLines = False
ws_kanban.column_dimensions["A"].width = 2
for col, w in zip("BDF", [46, 46, 46]):
    ws_kanban.column_dimensions[col].width = w
for col in "CE":
    ws_kanban.column_dimensions[col].width = 2

ws_kanban["B2"] = "🗂 칸반 보드"
ws_kanban["B2"].font = f(16, True)
ws_kanban["B3"] = "업무목록 시트에서 상태를 바꾸면 자동으로 갱신됩니다. (열별 최대 30개 표시 · 🔺 우선순위 높음 · ⚠ 마감 지연)"
ws_kanban["B3"].font = f(9, color=MUTED)

KB = [("B", "To Do", MUTED, GRAY_LT, BASE, INK),
      ("D", "Doing", BLUE, BLUE_LT, BLUE_MD, INK),
      ("F", "Done", GREEN, GREEN_LT, GREEN_MD, INK2)]
for col, st, head_fill, card_fill, card_bd, card_ink in KB:
    h = ws_kanban[f"{col}4"]
    h.value = f'="{st}  ("&COUNTIF(업무목록!$F$2:$F${LAST},"{st}")&"건)"'
    h.font = f(11, True, "FFFFFF")
    h.fill = fill(head_fill)
    h.alignment = Alignment(horizontal="center", vertical="center")
    ws_kanban.row_dimensions[4].height = 26

    helper = {"To Do": "L", "Doing": "M", "Done": "N"}[st]
    for k in range(1, KANBAN_SLOTS + 1):
        r = 4 + k
        cell = ws_kanban[f"{col}{r}"]
        cell.value = (f'=IFERROR(INDEX(업무목록!$O$2:$O${LAST},'
                      f'MATCH(ROW()-4,업무목록!${helper}$2:${helper}${LAST},0)),"")')
        cell.font = f(10, color=card_ink)
        cell.alignment = Alignment(vertical="center", indent=1)
        ws_kanban.row_dimensions[r].height = 24
    side = Side(style="thin", color=card_bd)
    ws_kanban.conditional_formatting.add(
        f"{col}5:{col}{4 + KANBAN_SLOTS}",
        FormulaRule(formula=[f'${col}5<>""'], fill=fill(card_fill),
                    border=Border(left=side, right=side, top=side, bottom=side)))

# ═══════════════════════════════ 대시보드 ═══════════════════════════════
ws_dash.sheet_view.showGridLines = False
paint(ws_dash, 1, 1, 36, 18, PAGE)
ws_dash.column_dimensions["A"].width = 2
for i in range(2, 14):
    ws_dash.column_dimensions[get_column_letter(i)].width = 11.5

ws_dash.merge_cells("B1:G1")
ws_dash["B1"] = "📊 PM 업무 대시보드"
ws_dash["B1"].font = f(18, True)
ws_dash["B1"].alignment = Alignment(vertical="center")
ws_dash.row_dimensions[1].height = 32
ws_dash["K1"] = "기준일"
ws_dash["K1"].font = f(9, color=MUTED)
ws_dash["K1"].alignment = Alignment(horizontal="right", vertical="center")
ws_dash.merge_cells("L1:M1")
ws_dash["L1"] = "=TODAY()"
ws_dash["L1"].number_format = "yyyy-mm-dd"
ws_dash["L1"].font = f(10, True, INK2)
ws_dash["L1"].alignment = Alignment(horizontal="center", vertical="center")

TOTAL = f'COUNTIF(업무목록!$B$2:$B${LAST},"<>")'
DONE = f'COUNTIF(업무목록!$F$2:$F${LAST},"Done")'
TILES = [
    (2, "전체 업무", f"={TOTAL}", "0", INK),
    (4, "To Do", f'=COUNTIF(업무목록!$F$2:$F${LAST},"To Do")', "0", INK2),
    (6, "Doing", f'=COUNTIF(업무목록!$F$2:$F${LAST},"Doing")', "0", BLUE),
    (8, "Done", f"={DONE}", "0", GREEN_TX),
    (10, "완료율", f"=IFERROR({DONE}/{TOTAL},0)", "0%", BLUE_DK),
    (12, "지연 업무", f'=COUNTIFS(업무목록!$B$2:$B${LAST},"<>",'
                     f'업무목록!$F$2:$F${LAST},"<>Done",'
                     f'업무목록!$H$2:$H${LAST},"<"&TODAY())', "0", RED),
]
ws_dash.row_dimensions[3].height = 20
ws_dash.row_dimensions[4].height = 38
for c0, label, formula, fmt, color in TILES:
    l1 = get_column_letter(c0)
    l2 = get_column_letter(c0 + 1)
    paint(ws_dash, 3, c0, 4, c0 + 1, SURFACE)
    ws_dash.merge_cells(f"{l1}3:{l2}3")
    ws_dash.merge_cells(f"{l1}4:{l2}4")
    lab = ws_dash[f"{l1}3"]
    lab.value = label
    lab.font = f(9, color=MUTED)
    lab.alignment = Alignment(horizontal="center", vertical="bottom")
    val = ws_dash[f"{l1}4"]
    val.value = formula
    val.number_format = fmt
    val.font = f(20, True, color)
    val.alignment = Alignment(horizontal="center", vertical="center")
    box(ws_dash, 3, c0, 4, c0 + 1, HAIR)

# ── 소계 테이블 ──
def table_header(ws, row, col, texts, widths=None):
    for i, t in enumerate(texts):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = f(9, True, INK2)
        c.fill = fill(GRAY_LT)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_bd

for col, title in [(2, "상태별 현황"), (5, "우선순위별 현황"), (8, "담당자별 현황")]:
    c = ws_dash.cell(row=7, column=col, value=title)
    c.font = f(12, True)

table_header(ws_dash, 8, 2, ["상태", "건수"])
for i, st in enumerate(STATUS):
    r = 9 + i
    ws_dash.cell(row=r, column=2, value=st)
    ws_dash.cell(row=r, column=3, value=f'=COUNTIF(업무목록!$F$2:$F${LAST},$B{r})')
    for c in (2, 3):
        cell = ws_dash.cell(row=r, column=c)
        cell.font = f(10)
        cell.fill = fill(SURFACE)
        cell.border = thin_bd
        cell.alignment = Alignment(horizontal="center", vertical="center")

table_header(ws_dash, 8, 5, ["우선순위", "건수"])
for i, p in enumerate(PRIORITY):
    r = 9 + i
    ws_dash.cell(row=r, column=5, value=p)
    ws_dash.cell(row=r, column=6, value=f'=COUNTIF(업무목록!$E$2:$E${LAST},$E{r})')
    for c in (5, 6):
        cell = ws_dash.cell(row=r, column=c)
        cell.font = f(10)
        cell.fill = fill(SURFACE)
        cell.border = thin_bd
        cell.alignment = Alignment(horizontal="center", vertical="center")

table_header(ws_dash, 8, 8, ["담당자", "To Do", "Doing", "Done", "합계", "완료율"])
for i in range(10):
    r = 9 + i
    ws_dash.cell(row=r, column=8, value=f'=IF(설정!$C{4 + i}="","",설정!$C{4 + i})')
    for j, st in enumerate(STATUS):
        ws_dash.cell(row=r, column=9 + j, value=(
            f'=IF($H{r}="","",COUNTIFS(업무목록!$D$2:$D${LAST},$H{r},'
            f'업무목록!$F$2:$F${LAST},"{st}"))'))
    ws_dash.cell(row=r, column=12, value=f'=IF($H{r}="","",$I{r}+$J{r}+$K{r})')
    ws_dash.cell(row=r, column=13, value=f'=IF($H{r}="","",IF($L{r}=0,"—",$K{r}/$L{r}))')
    ws_dash.cell(row=r, column=13).number_format = "0%"
    for c in range(8, 14):
        cell = ws_dash.cell(row=r, column=c)
        cell.font = f(10)
        cell.fill = fill(SURFACE)
        cell.border = thin_bd
        cell.alignment = Alignment(horizontal="center", vertical="center")

# ── 차트 ──
def val_labels():
    d = DataLabelList()
    d.showVal = True
    d.showSerName = False
    d.showCatName = False
    d.showLegendKey = False
    d.showPercent = False
    d.showBubbleSize = False
    return d

STATUS_COLORS = [MUTED, BLUE, GREEN]

bar = BarChart()
bar.type = "bar"
bar.title = "상태별 업무 현황"
data = Reference(ws_dash, min_col=3, min_row=8, max_row=11)
cats = Reference(ws_dash, min_col=2, min_row=9, max_row=11)
bar.add_data(data, titles_from_data=True)
bar.set_categories(cats)
bar.legend = None
bar.dataLabels = val_labels()
bar.gapWidth = 60
s = bar.series[0]
s.data_points = [DataPoint(idx=i, spPr=GraphicalProperties(solidFill=c))
                 for i, c in enumerate(STATUS_COLORS)]
bar.width = 8
bar.height = 6.5
ws_dash.add_chart(bar, "B14")

dough = DoughnutChart()
dough.title = "우선순위 분포"
data = Reference(ws_dash, min_col=6, min_row=8, max_row=11)
cats = Reference(ws_dash, min_col=5, min_row=9, max_row=11)
dough.add_data(data, titles_from_data=True)
dough.set_categories(cats)
dough.holeSize = 55
dough.dataLabels = val_labels()
s = dough.series[0]
s.data_points = [DataPoint(idx=i, spPr=GraphicalProperties(solidFill=c))
                 for i, c in enumerate([RED, AMBER, MUTED])]
dough.width = 8
dough.height = 6.5
ws_dash.add_chart(dough, "F14")

stack = BarChart()
stack.type = "bar"
stack.grouping = "stacked"
stack.overlap = 100
stack.title = "담당자별 업무 현황"
data = Reference(ws_dash, min_col=9, max_col=11, min_row=8, max_row=18)
cats = Reference(ws_dash, min_col=8, min_row=9, max_row=18)
stack.add_data(data, titles_from_data=True)
stack.set_categories(cats)
for srs, c in zip(stack.series, STATUS_COLORS):
    srs.graphicalProperties = GraphicalProperties(solidFill=c)
if stack.legend:
    stack.legend.position = "b"
stack.gapWidth = 40
stack.width = 12
stack.height = 6.5
ws_dash.add_chart(stack, "J14")

# ═══════════════════════════════ 사용법 ═══════════════════════════════
ws_help.sheet_view.showGridLines = False
ws_help.column_dimensions["A"].width = 3
ws_help.column_dimensions["B"].width = 110

LINES = [
    ("🗂 PM 업무관리 대시보드 — 사용법", 16, True, INK),
    ("To Do / Doing / Done 칸반 방식으로 업무를 관리하는 PM용 일정관리 툴입니다.", 10, False, INK2),
    ("", 10, False, INK),
    ("■ 시트 구성", 12, True, INK),
    ("   📊 대시보드 : 전체 현황 요약 — 자동 계산되므로 직접 입력하지 않습니다.", 10, False, INK),
    ("   🗂 칸반보드 : To Do / Doing / Done 보드 — 자동 갱신되므로 직접 입력하지 않습니다.", 10, False, INK),
    ("   ✏️ 업무목록 : ★ 실제로 입력하는 시트입니다. 업무 추가/수정은 모두 여기서 하세요.", 10, True, GREEN_TX),
    ("   ⚙️ 설정 : 담당자 드롭다운 목록을 관리합니다.", 10, False, INK),
    ("", 10, False, INK),
    ("■ 업무 추가하기 (업무목록 시트)", 12, True, INK),
    ("   1) 비어 있는 행에 '업무명'부터 입력하세요.", 10, False, INK),
    ("   2) 담당자 / 우선순위 / 상태는 셀을 클릭하면 나오는 드롭다운(▼)에서 선택하세요.", 10, False, INK),
    ("   3) 시작일 / 마감일은 날짜로 입력하세요. 예: 2026-07-10", 10, False, INK),
    ("   4) 진행률은 0~100 사이 숫자로 입력하세요.", 10, False, INK),
    ("   5) No. 열과 D-Day 열은 자동 계산되므로 입력하지 마세요.", 10, False, INK),
    ("", 10, False, INK),
    ("■ 업무 상태 바꾸기", 12, True, INK),
    ("   '상태' 열 드롭다운에서 To Do → Doing → Done 으로 바꾸면 칸반보드와 대시보드에 즉시 반영됩니다.", 10, False, INK),
    ("", 10, False, INK),
    ("■ 자동 표시 규칙", 12, True, INK),
    ("   · 마감일이 지났는데 Done이 아닌 업무 → 행 전체가 붉게 표시되고 D-Day에 '⚠ n일 지연'", 10, False, INK),
    ("   · 오늘이 마감일 → '🔔 오늘 마감'", 10, False, INK),
    ("   · 완료(Done)된 업무 → 글자가 회색 처리", 10, False, INK),
    ("   · 우선순위 '높음' → 빨간색 강조, 칸반 카드에 🔺 표시", 10, False, INK),
    ("", 10, False, INK),
    ("■ 담당자 목록 수정 (설정 시트)", 12, True, INK),
    ("   초록색 담당자 칸(C4~C13)에서 이름을 바꾸거나 추가하세요. 최대 10명까지 가능합니다.", 10, False, INK),
    ("   상태(To Do/Doing/Done)와 우선순위(높음/중간/낮음) 목록은 수식과 연결되어 있으니 바꾸지 마세요.", 10, False, RED),
    ("", 10, False, INK),
    ("■ 알아두면 좋은 점", 12, True, INK),
    (f"   · 업무는 {N_ROWS}행까지 서식과 자동 수식이 미리 적용되어 있습니다. 더 필요하면 마지막 행을 복사해 아래로 붙여넣으세요.", 10, False, INK),
    (f"   · 칸반보드는 각 열에서 위에서부터 {KANBAN_SLOTS}개까지 표시합니다.", 10, False, INK),
    ("   · D-Day와 지연 표시는 파일을 여는 시점의 '오늘' 날짜 기준으로 자동 계산됩니다.", 10, False, INK),
    ("   · 처음에 들어 있는 예시 업무 12건은 자유롭게 지우고 사용하세요.", 10, False, INK),
]
for i, (text, size, bold, color) in enumerate(LINES, start=2):
    c = ws_help.cell(row=i, column=2, value=text)
    c.font = f(size, bold, color)
    c.alignment = Alignment(vertical="center")
    ws_help.row_dimensions[i].height = 24 if size >= 12 else 18

wb.active = wb.index(ws_dash)
wb.properties.title = "PM 업무관리 대시보드"
wb.properties.description = "To Do / Doing / Done 칸반 방식 일정관리 툴"

wb.save(OUT)
print(f"저장 완료: {OUT}")
